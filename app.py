import streamlit as st
import os
import re
import shutil
import requests
from datetime import datetime, timedelta, timezone
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image as PILImage, ExifTags
from io import BytesIO

# ==========================================
# PAGE CONFIG
# ==========================================
st.set_page_config(page_title="QC Master Pro - LGJA", page_icon="QC", layout="wide")

# ==========================================
# FUNCTIONS
# ==========================================

def log_traffic(user, toko, tgl_qc, layout):
    JSONBLOB_ID = "019e8740-72c4-7731-8328-0e2c67465233"
    API_URL = f"https://jsonblob.com/api/jsonBlob/{JSONBLOB_ID}"
    try:
        response = requests.get(API_URL)
        try:
            data = response.json()
        except:
            data = []
        if not isinstance(data, list):
            data = []
        tz_jkt = timezone(timedelta(hours=7))
        timestamp = datetime.now(tz_jkt).strftime("%d %B %y / %H:%M")
        data.append({
            "Nama Pengguna": user,
            "Nama Toko": toko,
            "Tanggal QC": tgl_qc,
            "Timestamp": timestamp,
            "Ukuran Layout": layout,
        })
        headers = {"Content-Type": "application/json", "Accept": "application/json"}
        put_response = requests.put(API_URL, json=data, headers=headers)
        if put_response.status_code in [200, 201]:
            st.toast("Traffic berhasil dicatat di server!")
        else:
            st.error(f"Gagal mencatat traffic! Status: {put_response.status_code}")
    except Exception as e:
        st.error(f"Sistem Traffic Error Koneksi: {e}")


def correct_orientation(img):
    try:
        if hasattr(img, '_getexif') and img._getexif() is not None:
            exif = img._getexif()
            orientation = exif.get(ExifTags.Base.Orientation)
            if orientation == 3:
                img = img.rotate(180, expand=True)
            elif orientation == 6:
                img = img.rotate(270, expand=True)
            elif orientation == 8:
                img = img.rotate(90, expand=True)
    except:
        pass
    return img


def extract_datetime(filename, uploaded_file_obj):
    match = re.search(r'(\d{4}-\d{2}-\d{2}) at (\d{2}\.\d{2}\.\d{2})', filename)
    if match:
        try:
            return datetime.strptime(
                f"{match.group(1)} {match.group(2).replace('.', ':')}", "%Y-%m-%d %H:%M:%S"
            )
        except:
            pass
    return datetime.min


# ==========================================
# CSS - iLOVEPDF DARK THEME
# ==========================================
st.markdown(
    """
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&display=swap" rel="stylesheet">

    <style>
    /* ========== BASE ========== */
    * { box-sizing: border-box; }

    html, body,
    [data-testid="stAppViewContainer"],
    [data-testid="stHeader"] {
        background: #0b0b14 !important;
        color: #d0d0e0 !important;
        font-family: 'Plus Jakarta Sans', -apple-system, sans-serif !important;
    }

    .main .block-container {
        max-width: 960px !important;
        padding: 0 1rem 3rem !important;
    }

    header[data-testid="stHeader"] { display: none !important; }
    #MainMenu { visibility: hidden; }
    footer { visibility: hidden; }

    ::-webkit-scrollbar { width: 6px; }
    ::-webkit-scrollbar-track { background: #0b0b14; }
    ::-webkit-scrollbar-thumb { background: #252540; border-radius: 3px; }
    ::-webkit-scrollbar-thumb:hover { background: #35355a; }

    /* ========== TOP NAV ========== */
    .topnav {
        position: sticky;
        top: 0;
        z-index: 999;
        display: flex;
        align-items: center;
        justify-content: space-between;
        margin: 0 -1rem;
        padding: 0.85rem 2rem;
        background: rgba(11, 11, 20, 0.88);
        backdrop-filter: blur(24px);
        -webkit-backdrop-filter: blur(24px);
        border-bottom: 1px solid rgba(255,255,255,0.05);
    }

    .topnav .logo {
        font-size: 1.2rem;
        font-weight: 800;
        letter-spacing: -0.03em;
        color: #fff;
    }
    .topnav .logo em {
        font-style: normal;
        color: #ff4655;
    }

    .topnav .nav-right {
        display: flex;
        align-items: center;
        gap: 1rem;
    }
    .topnav .nav-pill {
        padding: 0.3rem 0.8rem;
        border-radius: 20px;
        background: rgba(255,70,85,0.1);
        color: #ff6b7a;
        font-size: 0.72rem;
        font-weight: 700;
        letter-spacing: 0.04em;
    }
    .topnav .nav-author {
        color: #55556a;
        font-size: 0.78rem;
        font-weight: 500;
    }

    /* ========== HERO ========== */
    .hero {
        position: relative;
        text-align: center;
        padding: 3.5rem 1rem 1.5rem;
    }
    .hero::before {
        content: '';
        position: absolute;
        top: 10px;
        left: 50%;
        transform: translateX(-50%);
        width: 500px;
        height: 260px;
        background: radial-gradient(ellipse, rgba(255,70,85,0.1) 0%, transparent 70%);
        pointer-events: none;
        z-index: 0;
    }
    .hero h1 {
        position: relative;
        z-index: 1;
        font-size: 2.6rem;
        font-weight: 800;
        letter-spacing: -0.04em;
        line-height: 1.1;
        background: linear-gradient(135deg, #ff4655 0%, #ff6b81 50%, #fff 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
        margin-bottom: 0.6rem;
    }
    .hero p {
        position: relative;
        z-index: 1;
        color: #5a5a78;
        font-size: 1rem;
        font-weight: 400;
    }

    /* ========== STEPS GRID ========== */
    .steps-grid {
        display: grid;
        grid-template-columns: repeat(4, 1fr);
        gap: 0.75rem;
        margin: 1.5rem 0 2rem;
    }

    .s-card {
        background: #12122a;
        border: 1px solid rgba(255,255,255,0.05);
        border-radius: 14px;
        padding: 1.2rem 0.8rem;
        text-align: center;
        transition: all 0.3s ease;
    }
    .s-card:hover {
        border-color: rgba(255,70,85,0.35);
        transform: translateY(-3px);
        box-shadow: 0 12px 28px rgba(255,70,85,0.08);
    }

    .s-card .s-num {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        width: 38px;
        height: 38px;
        border-radius: 12px;
        background: linear-gradient(135deg, #ff4655, #c0354a);
        color: #fff;
        font-weight: 800;
        font-size: 0.9rem;
        margin-bottom: 0.6rem;
    }
    .s-card .s-title {
        font-weight: 700;
        font-size: 0.88rem;
        color: #e0e0f0;
        margin-bottom: 0.25rem;
    }
    .s-card .s-desc {
        color: #505068;
        font-size: 0.72rem;
        line-height: 1.45;
    }

    /* ========== SECTION HEADERS ========== */
    .sec-head {
        display: flex;
        align-items: center;
        gap: 0.75rem;
        margin-bottom: 1rem;
        padding-bottom: 0.8rem;
        border-bottom: 1px solid rgba(255,255,255,0.05);
    }
    .sec-head .sh-num {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        min-width: 36px;
        height: 36px;
        border-radius: 10px;
        background: linear-gradient(135deg, #ff4655, #c0354a);
        color: #fff;
        font-weight: 800;
        font-size: 0.82rem;
    }
    .sec-head .sh-title {
        font-weight: 700;
        font-size: 1.05rem;
        color: #e8e8f8;
    }
    .sec-head .sh-sub {
        font-size: 0.76rem;
        color: #505068;
        font-weight: 400;
    }

    /* ========== CARDS (bordered containers) ========== */
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: #12122a !important;
        border: 1px solid rgba(255,255,255,0.06) !important;
        border-radius: 16px !important;
        margin-bottom: 0.9rem !important;
        overflow: hidden !important;
        transition: border-color 0.3s ease, box-shadow 0.3s ease !important;
    }
    [data-testid="stVerticalBlockBorderWrapper"]:hover {
        border-color: rgba(255,70,85,0.2) !important;
        box-shadow: 0 8px 32px rgba(0,0,0,0.25) !important;
    }

    /* ========== INPUTS ========== */
    [data-testid="stTextInput"] input,
    [data-testid="stNumberInput"] input {
        background: #0d0d1e !important;
        border: 1px solid rgba(255,255,255,0.08) !important;
        border-radius: 10px !important;
        color: #e8e8f0 !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-size: 0.88rem !important;
        padding: 0.65rem 0.9rem !important;
        transition: border-color 0.2s, box-shadow 0.2s !important;
    }
    [data-testid="stTextInput"] input:focus,
    [data-testid="stNumberInput"] input:focus {
        border-color: #ff4655 !important;
        box-shadow: 0 0 0 3px rgba(255,70,85,0.12) !important;
    }
    [data-testid="stTextInput"] input::placeholder {
        color: #3a3a55 !important;
    }

    [data-testid="stWidgetLabel"] label {
        color: #8888a5 !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 600 !important;
        font-size: 0.82rem !important;
    }

    /* ========== SELECTBOX ========== */
    [data-testid="stSelectbox"] > div > div {
        background: #0d0d1e !important;
        border: 1px solid rgba(255,255,255,0.08) !important;
        border-radius: 10px !important;
        color: #e8e8f0 !important;
    }
    [data-testid="stSelectbox"] > div > div:hover {
        border-color: rgba(255,70,85,0.4) !important;
    }

    /* ========== RADIO ========== */
    [data-testid="stRadio"] label {
        background: #0d0d1e !important;
        border: 1px solid rgba(255,255,255,0.06) !important;
        border-radius: 8px !important;
        padding: 0.4rem 0.85rem !important;
        transition: all 0.2s !important;
    }
    [data-testid="stRadio"] label:hover {
        border-color: #ff4655 !important;
        background: #1a1a30 !important;
    }

    /* ========== FILE UPLOADER ========== */
    [data-testid="stFileUploader"] {
        border: 2px dashed rgba(255,70,85,0.2) !important;
        border-radius: 14px !important;
        background: rgba(255,70,85,0.02) !important;
        padding: 1rem !important;
        transition: all 0.3s !important;
    }
    [data-testid="stFileUploader"]:hover {
        border-color: rgba(255,70,85,0.45) !important;
        background: rgba(255,70,85,0.05) !important;
    }
    [data-testid="stFileUploader"] button {
        background: #1a1a35 !important;
        color: #ff6b7a !important;
        border: 1px solid rgba(255,70,85,0.2) !important;
        border-radius: 8px !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 600 !important;
    }
    [data-testid="stFileUploader"] button:hover {
        background: #ff4655 !important;
        color: #fff !important;
        border-color: #ff4655 !important;
    }

    /* ========== BUTTONS ========== */
    [data-testid="stButton"] button[kind="primary"] {
        background: linear-gradient(135deg, #ff4655, #d03040) !important;
        color: #fff !important;
        border: none !important;
        border-radius: 12px !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 700 !important;
        font-size: 1rem !important;
        padding: 0.8rem 2rem !important;
        letter-spacing: 0.03em;
        transition: all 0.3s !important;
        box-shadow: 0 4px 20px rgba(255,70,85,0.3) !important;
    }
    [data-testid="stButton"] button[kind="primary"]:hover {
        background: linear-gradient(135deg, #ff5a68, #ff4655) !important;
        box-shadow: 0 6px 28px rgba(255,70,85,0.45) !important;
        transform: translateY(-1px);
    }

    [data-testid="stButton"] button:not([kind="primary"]) {
        background: #1a1a30 !important;
        color: #ff6b7a !important;
        border: 1px solid rgba(255,70,85,0.15) !important;
        border-radius: 10px !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 600 !important;
        transition: all 0.2s !important;
    }
    [data-testid="stButton"] button:not([kind="primary"]):hover {
        background: #20203a !important;
        border-color: #ff4655 !important;
    }

    /* ========== DOWNLOAD BUTTON ========== */
    [data-testid="stDownloadButton"] a,
    [data-testid="stDownloadButton"] button {
        background: linear-gradient(135deg, #ff4655, #d03040) !important;
        color: #fff !important;
        border: none !important;
        border-radius: 12px !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 700 !important;
        font-size: 1rem !important;
        padding: 0.8rem 2rem !important;
        transition: all 0.3s !important;
        animation: pulseRed 2s ease-in-out infinite !important;
    }
    @keyframes pulseRed {
        0%, 100% { box-shadow: 0 4px 20px rgba(255,70,85,0.3); }
        50%      { box-shadow: 0 4px 32px rgba(255,70,85,0.55); }
    }
    [data-testid="stDownloadButton"] a:hover,
    [data-testid="stDownloadButton"] button:hover {
        box-shadow: 0 6px 28px rgba(255,70,85,0.5) !important;
        transform: translateY(-1px);
    }

    /* ========== ALERTS ========== */
    [data-testid="stAlert"] {
        border-radius: 10px !important;
        border: none !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-size: 0.85rem !important;
    }
    [data-testid="stAlert"][kind="success"] {
        background: rgba(16,185,129,0.1) !important;
        color: #6ee7b7 !important;
        border-left: 3px solid #10b981 !important;
    }
    [data-testid="stAlert"][kind="warning"] {
        background: rgba(245,158,11,0.1) !important;
        color: #fcd34d !important;
        border-left: 3px solid #f59e0b !important;
    }
    [data-testid="stAlert"][kind="error"] {
        background: rgba(239,68,68,0.1) !important;
        color: #fca5a5 !important;
        border-left: 3px solid #ef4444 !important;
    }
    [data-testid="stAlert"][kind="info"] {
        background: rgba(255,70,85,0.06) !important;
        color: #ff9ba5 !important;
        border-left: 3px solid #ff4655 !important;
    }

    [data-testid="stToast"] {
        background: #1a1a30 !important;
        color: #ff6b7a !important;
        border: 1px solid rgba(255,70,85,0.2) !important;
        border-radius: 10px !important;
    }

    /* ========== MISC ========== */
    [data-testid="stSpinner"] { color: #ff4655 !important; }
    [data-testid="stHorizontalBlock"] { gap: 1rem !important; }

    [data-testid="stNumberInput"] button {
        background: #1a1a35 !important;
        color: #ff6b7a !important;
        border: 1px solid rgba(255,70,85,0.15) !important;
    }
    [data-testid="stNumberInput"] button:hover {
        background: #ff4655 !important;
        color: #fff !important;
    }

    hr {
        border: none !important;
        height: 1px !important;
        background: linear-gradient(90deg, transparent, rgba(255,255,255,0.06), transparent) !important;
        margin: 1.5rem 0 !important;
    }

    [data-testid="stMarkdownContainer"] h3 {
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 700 !important;
        font-size: 1.1rem !important;
        color: #e0e0f5 !important;
    }
    [data-testid="stMarkdownContainer"] h5 {
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        font-weight: 600 !important;
        color: #7a7a95 !important;
        font-size: 0.88rem !important;
    }
    [data-testid="stCaption"] {
        color: #4a4a65 !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
    }

    /* ========== FOOTER ========== */
    .app-footer {
        text-align: center;
        padding: 2.5rem 1rem 1rem;
        margin-top: 2rem;
        border-top: 1px solid rgba(255,255,255,0.04);
    }
    .app-footer .f-logo {
        font-weight: 800;
        font-size: 0.95rem;
        color: #ff4655;
        margin-bottom: 0.35rem;
    }
    .app-footer .f-text {
        color: #35354a;
        font-size: 0.73rem;
    }

    /* ========== ANIMATIONS ========== */
    @keyframes fadeUp {
        from { opacity: 0; transform: translateY(18px); }
        to   { opacity: 1; transform: translateY(0); }
    }

    .hero      { animation: fadeUp 0.6s ease forwards; }
    .steps-grid { animation: fadeUp 0.6s ease 0.12s forwards; opacity: 0; }

    [data-testid="stVerticalBlock"] > div {
        animation: fadeUp 0.5s ease forwards;
        opacity: 0;
    }
    [data-testid="stVerticalBlock"] > div:nth-child(1)  { animation-delay: 0.02s; }
    [data-testid="stVerticalBlock"] > div:nth-child(2)  { animation-delay: 0.07s; }
    [data-testid="stVerticalBlock"] > div:nth-child(3)  { animation-delay: 0.12s; }
    [data-testid="stVerticalBlock"] > div:nth-child(4)  { animation-delay: 0.17s; }
    [data-testid="stVerticalBlock"] > div:nth-child(5)  { animation-delay: 0.22s; }
    [data-testid="stVerticalBlock"] > div:nth-child(6)  { animation-delay: 0.27s; }
    [data-testid="stVerticalBlock"] > div:nth-child(7)  { animation-delay: 0.32s; }
    [data-testid="stVerticalBlock"] > div:nth-child(8)  { animation-delay: 0.37s; }
    [data-testid="stVerticalBlock"] > div:nth-child(9)  { animation-delay: 0.42s; }
    [data-testid="stVerticalBlock"] > div:nth-child(10) { animation-delay: 0.47s; }
    [data-testid="stVerticalBlock"] > div:nth-child(11) { animation-delay: 0.52s; }
    [data-testid="stVerticalBlock"] > div:nth-child(12) { animation-delay: 0.57s; }
    [data-testid="stVerticalBlock"] > div:nth-child(13) { animation-delay: 0.62s; }
    [data-testid="stVerticalBlock"] > div:nth-child(14) { animation-delay: 0.67s; }
    [data-testid="stVerticalBlock"] > div:nth-child(15) { animation-delay: 0.72s; }

    /* ========== RESPONSIVE ========== */
    @media (max-width: 768px) {
        .steps-grid { grid-template-columns: repeat(2, 1fr); }
        .hero h1   { font-size: 1.8rem; }
        .topnav     { padding: 0.7rem 1rem; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ==========================================
# NAVIGATION BAR
# ==========================================
st.markdown(
    """
    <div class="topnav">
        <div class="logo"><em>QC</em> Master Pro</div>
        <div class="nav-right">
            <span class="nav-pill">v2.0 LGJA</span>
            <span class="nav-author">Build by GJorma</span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ==========================================
# HERO SECTION
# ==========================================
st.markdown(
    """
    <div class="hero">
        <h1>QC Master Pro</h1>
        <p>Automated photo arrangement & compression for quality control reports</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ==========================================
# STEPS OVERVIEW (Decorative)
# ==========================================
st.markdown(
    """
    <div class="steps-grid">
        <div class="s-card">
            <div class="s-num">1</div>
            <div class="s-title">Info Lokasi</div>
            <div class="s-desc">Tentukan nama toko dan tanggal QC</div>
        </div>
        <div class="s-card">
            <div class="s-num">2</div>
            <div class="s-title">Template</div>
            <div class="s-desc">Upload atau pilih Excel master</div>
        </div>
        <div class="s-card">
            <div class="s-num">3</div>
            <div class="s-title">Foto QC</div>
            <div class="s-desc">Upload foto dari lapangan</div>
        </div>
        <div class="s-card">
            <div class="s-num">4</div>
            <div class="s-title">Export</div>
            <div class="s-desc">Download hasil proses</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ==========================================
# USER NAME (Gatekeeper)
# ==========================================
st.markdown(
    """
    <div class="sec-head" style="margin-top:0.3rem">
        <div class="sh-num">U</div>
        <div>
            <div class="sh-title">Identitas Pengguna</div>
            <div class="sh-sub">Wajib diisi untuk memulai proses</div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

user_name = st.text_input(
    "Nama Pengguna", placeholder="Masukkan nama Anda...", label_visibility="collapsed"
)
if not user_name:
    st.warning("Silakan isi Nama Pengguna di atas agar sistem bisa digunakan.")
    st.stop()


# ==========================================
# SECTION 1: INFO LOKASI
# ==========================================
with st.container(border=True):
    st.markdown(
        """
        <div class="sec-head">
            <div class="sh-num">1</div>
            <div>
                <div class="sh-title">Informasi Lokasi QC</div>
                <div class="sh-sub">Detail lokasi untuk nama file output</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    col1, col2 = st.columns(2)
    with col1:
        nama_toko = st.text_input("Nama Toko / Area", placeholder="Contoh: Batavia PIK")
    with col2:
        tanggal_qc = st.text_input("Tanggal QC", placeholder="Contoh: 10 Mar 26")


# ==========================================
# SECTION 2: TEMPLATE EXCEL
# ==========================================
with st.container(border=True):
    st.markdown(
        """
        <div class="sec-head">
            <div class="sh-num">2</div>
            <div>
                <div class="sh-title">Template Excel Master</div>
                <div class="sh-sub">Upload manual atau pilih preset dari server</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    mode_template = st.radio(
        "Sumber template:", ["Upload Manual", "File Preset (GitHub)"], horizontal=True
    )

    excel_file = None
    if mode_template == "Upload Manual":
        excel_file = st.file_uploader("Pilih file Excel (.xlsx)", type=["xlsx"])
    else:
        preset_pilihan = st.selectbox(
            "Pilih file preset:", ["QC_LGJA.xlsx", "QC_Sultan.xlsx", "QC_Vano.xlsx"]
        )
        if os.path.exists(preset_pilihan):
            with open(preset_pilihan, "rb") as f:
                excel_file = BytesIO(f.read())
        else:
            st.error(f"File '{preset_pilihan}' belum tersedia di folder server/GitHub Anda.")

    selected_sheet = None
    if excel_file:
        try:
            excel_file.seek(0)
            wb_scan = load_workbook(excel_file, read_only=True)
            sheet_names = wb_scan.sheetnames
            wb_scan.close()
            selected_sheet = st.selectbox("Target Sheet:", sheet_names)
        except Exception as e:
            st.error(f"Gagal membaca struktur Excel: {e}")


# ==========================================
# SECTION 3: LAYOUT SETTINGS
# ==========================================
with st.container(border=True):
    st.markdown(
        """
        <div class="sec-head">
            <div class="sh-num">3</div>
            <div>
                <div class="sh-title">Pengaturan Layout & Ukuran</div>
                <div class="sh-sub">Konfigurasi tata letak gambar di dalam Excel</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    layout_option = st.selectbox("Pilih Opsi Layout:", ["LGJA", "Sultan", "Vano", "Custom"])

    if layout_option == "LGJA":
        ROWS = [2, 4, 6, 8, 10, 12]
        COLS = list(range(1, 13))
        COL_W = 41
        ROW_H = 246
        IMAGE_WIDTH_CM = 6.4
        IMAGE_HEIGHT_CM = 8.30
    elif layout_option in ["Sultan", "Vano"]:
        ROWS = [2, 4, 6, 8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30]
        COLS = [1, 2, 3, 4, 5, 6]
        COL_W = 20.43
        ROW_H = 123.75
        IMAGE_WIDTH_CM = 3.2
        IMAGE_HEIGHT_CM = 4.10
    else:
        c_row, c_col = st.columns(2)
        r_in = c_row.text_input("Rows (pisahkan koma)", "2,4,6,8,10,12")
        c_in = c_col.text_input("Columns (1=A, 2=B, koma)", "1,2,3,4,5,6")
        try:
            ROWS = [int(x.strip()) for x in r_in.split(",")]
            COLS = [int(x.strip()) for x in c_in.split(",")]
        except:
            st.error("Format Rows/Columns salah!")
            st.stop()
        c_w, c_h = st.columns(2)
        COL_W = c_w.number_input("Col Width", value=20.43)
        ROW_H = c_h.number_input("Row Height", value=123.75)
        i_w, i_h = st.columns(2)
        IMAGE_WIDTH_CM = i_w.number_input("Image Width (cm)", value=3.2)
        IMAGE_HEIGHT_CM = i_h.number_input("Image Height (cm)", value=4.10)


# ==========================================
# SECTION 4: UPLOAD FOTO
# ==========================================
with st.container(border=True):
    st.markdown(
        """
        <div class="sec-head">
            <div class="sh-num">4</div>
            <div>
                <div class="sh-title">Upload Foto QC Lapangan</div>
                <div class="sh-sub">Pilih banyak foto sekaligus dari galeri HP</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.info("Tekan tahan untuk memilih banyak foto sekaligus dari galeri HP Anda.")

    if "uploader_key" not in st.session_state:
        st.session_state.uploader_key = 0

    if st.button("Hapus Semua Foto (Reset)"):
        st.session_state.uploader_key += 1
        st.rerun()

    uploaded_photos = st.file_uploader(
        "Pilih semua foto sekaligus dari Galeri",
        type=["jpg", "jpeg", "png", "webp"],
        accept_multiple_files=True,
        key=f"photos_{st.session_state.uploader_key}",
    )

    if uploaded_photos:
        st.success(f"Photo : {len(uploaded_photos)} foto dipilih")
    else:
        st.caption("Belum ada foto yang dipilih")


# ==========================================
# SECTION 5: EKSEKUSI
# ==========================================
with st.container(border=True):
    st.markdown(
        """
        <div class="sec-head">
            <div class="sh-num">
                <svg width="16" height="16" viewBox="0 0 24 24" fill="none"
                     stroke="white" stroke-width="2.5" stroke-linecap="round"
                     stroke-linejoin="round">
                    <polygon points="5 3 19 12 5 21 5 3"></polygon>
                </svg>
            </div>
            <div>
                <div class="sh-title">Eksekusi Proses Data</div>
                <div class="sh-sub">Klik tombol di bawah untuk memulai</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if st.button("MULAI EXPORT DAN PROSES DATA", type="primary", use_container_width=True):
        if not nama_toko or not tanggal_qc:
            st.warning("Silakan isi Nama Toko dan Tanggal QC!")
        elif not excel_file:
            st.warning("Silakan upload atau pilih Excel Template Master!")
        elif not uploaded_photos:
            st.warning("Silakan pilih foto QC yang akan diproses!")
        elif not selected_sheet:
            st.error("Target sheet tidak valid.")
        else:
            with st.spinner("Sedang memproses dan mengompres foto... Mohon tunggu..."):
                try:
                    log_traffic(user_name, nama_toko, tanggal_qc, layout_option)

                    excel_file.seek(0)
                    wb = load_workbook(excel_file)
                    ws = wb[selected_sheet]

                    for c in COLS:
                        ws.column_dimensions[chr(64 + c)].width = COL_W
                    for r in ROWS:
                        ws.row_dimensions[r].height = ROW_H

                    sorted_photos = sorted(
                        uploaded_photos,
                        key=lambda x: extract_datetime(x.name, x),
                        reverse=True,
                    )

                    all_cells = [
                        f"{chr(64 + col)}{row}" for row in ROWS for col in COLS
                    ]
                    success_count = 0

                    temp_dir = "temp_web_photos"
                    os.makedirs(temp_dir, exist_ok=True)

                    for i in range(min(len(sorted_photos), len(all_cells))):
                        photo = sorted_photos[i]
                        temp_path = os.path.join(temp_dir, f"compressed_img_{i}.jpg")

                        with PILImage.open(photo) as img_pil:
                            img_pil = correct_orientation(img_pil)
                            if img_pil.mode in ("RGBA", "P"):
                                img_pil = img_pil.convert("RGB")
                            img_pil.thumbnail(
                                (1280, 1280), PILImage.Resampling.LANCZOS
                            )
                            img_pil.save(
                                temp_path,
                                format="JPEG",
                                quality=82,
                                optimize=True,
                                subsampling=0,
                            )

                        img_excel = ExcelImage(temp_path)
                        img_excel.width = int(IMAGE_WIDTH_CM * 37.8)
                        img_excel.height = int(IMAGE_HEIGHT_CM * 37.8)
                        ws.add_image(img_excel, all_cells[i])
                        success_count += 1

                    output = BytesIO()
                    wb.save(output)
                    wb.close()
                    output.seek(0)

                    shutil.rmtree(temp_dir, ignore_errors=True)

                    final_filename = f"{nama_toko.strip()} {tanggal_qc.strip()}.xlsx"

                    st.success(
                        f"Berhasil menyusun & mengompres {success_count} foto! File siap diunduh."
                    )

                    st.download_button(
                        label="DOWNLOAD FILE",
                        data=output,
                        file_name=final_filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True,
                    )

                except Exception as e:
                    st.error(f"Terjadi kesalahan sistem saat memproses Excel: {e}")


# ==========================================
# FOOTER
# ==========================================
st.markdown(
    """
    <div class="app-footer">
        <div class="f-logo">QC Master Pro</div>
        <div class="f-text">Advanced Dynamic Automation System &mdash; Build by AI &amp; GJorma</div>
    </div>
    """,
    unsafe_allow_html=True,
)
